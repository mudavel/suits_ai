"""Read-only historical catalogue, joined by CNJ from the supplied workbook.

The in-memory index is built lazily and replaced when the workbook changes.
It is independent of operational cases and never modifies the source file.
"""
from pathlib import Path
import sqlite3
from threading import Lock
import unicodedata

from openpyxl import load_workbook

RESULT_SHEET = "Resultados dos processos"
SUBSIDY_SHEET = "Subsídios disponibilizados"
RESULT_COLUMNS = {
    "Número do processo": "case_number", "UF": "uf", "Assunto": "subject",
    "Sub-assunto": "sub_subject", "Resultado macro": "macro_result",
    "Resultado micro": "micro_result", "Valor da causa": "cause_value",
    "Valor da condenação/indenização": "condemnation_value",
}
SUBSIDY_COLUMNS = {
    "Contrato": "contract", "Extrato": "statement", "Comprovante de crédito": "credit_receipt",
    "Dossiê": "dossier", "Demonstrativo de evolução da dívida": "debt_evolution",
    "Laudo referenciado": "referenced_report",
}
SORT_COLUMNS = set(RESULT_COLUMNS.values()) | set(SUBSIDY_COLUMNS.values()) | {"source_row", "subsidy_count"}


def normalize(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value or '').casefold()) if not unicodedata.combining(c))


class HistoricalBase:
    def __init__(self, source: Path):
        self.source = source
        self._lock = Lock()
        self._db = None
        self._stamp = None
        self._options = None

    def close(self):
        with self._lock:
            if self._db is not None:
                self._db.close()
                self._db = None
                self._stamp = None

    def _load(self):
        stat = self.source.stat()
        stamp = (stat.st_mtime_ns, stat.st_size)
        if self._db is not None and self._stamp == stamp:
            return
        workbook = load_workbook(self.source, read_only=True, data_only=True)
        db = sqlite3.connect(':memory:', check_same_thread=False)
        db.row_factory = sqlite3.Row
        try:
            subsidy_rows = workbook[SUBSIDY_SHEET].iter_rows(min_row=2, values_only=True)
            headers = next(subsidy_rows)
            key = headers.index('Número do processos')
            positions = [headers.index(label) for label in SUBSIDY_COLUMNS]
            subsidies = {}
            for row_number, row in enumerate(subsidy_rows, 3):
                number = row[key]
                values = [row[index] for index in positions]
                if not number or number in subsidies or any(value not in (0, 1) for value in values):
                    raise ValueError('Identificador ou presença documental inválida na base histórica.')
                subsidies[number] = (row_number, *[int(value) for value in values])
            db.execute('''CREATE TABLE history (
                case_number TEXT PRIMARY KEY, uf TEXT, subject TEXT, sub_subject TEXT,
                macro_result TEXT, micro_result TEXT, cause_value REAL, condemnation_value REAL,
                source_row INTEGER, subsidy_row INTEGER, contract INTEGER, statement INTEGER,
                credit_receipt INTEGER, dossier INTEGER, debt_evolution INTEGER, referenced_report INTEGER,
                subsidy_count INTEGER, search_text TEXT)''')
            rows = workbook[RESULT_SHEET].iter_rows(values_only=True)
            headers = next(rows)
            positions = [headers.index(label) for label in RESULT_COLUMNS]
            records = []
            for row_number, row in enumerate(rows, 2):
                values = [row[index] for index in positions]
                number = values[0]
                if not number or number not in subsidies:
                    raise ValueError('Processo ausente ou sem correspondência na aba de subsídios.')
                presence = subsidies.pop(number)
                search_text = normalize(' '.join(str(value or '') for value in values[:6]))
                records.append((*values, row_number, *presence, sum(presence[1:]), search_text))
            if subsidies or not records:
                raise ValueError('As abas históricas devem conter os mesmos processos únicos.')
            db.executemany('INSERT INTO history VALUES (' + ','.join(['?'] * 18) + ')', records)
            db.execute('CREATE INDEX history_source ON history(source_row)')
            db.execute('CREATE INDEX history_filters ON history(uf, micro_result)')
            options = {field: [row[0] for row in db.execute(f'SELECT DISTINCT {field} FROM history WHERE {field} IS NOT NULL ORDER BY {field}')]
                       for field in ('uf', 'subject', 'sub_subject', 'macro_result', 'micro_result')}
            total = db.execute('SELECT count(*) FROM history').fetchone()[0]
        except Exception:
            db.close()
            raise
        finally:
            workbook.close()
        if self._db is not None:
            self._db.close()
        self._db, self._stamp, self._options, self._total = db, stamp, options, total

    def query(self, *, page=1, page_size=50, q='', uf='', subject='', sub_subject='',
              macro_result='', micro_result='', subsidy='', presence='provided', sort='source_row', direction='asc'):
        if sort not in SORT_COLUMNS or direction not in {'asc', 'desc'}:
            raise ValueError('Ordenação inválida.')
        if subsidy and subsidy not in SUBSIDY_COLUMNS.values():
            raise ValueError('Subsídio inválido.')
        clauses, params = [], []
        if q.strip():
            clauses.append('instr(search_text, ?) > 0')
            params.append(normalize(q.strip()))
        for field, value in [('uf', uf), ('subject', subject), ('sub_subject', sub_subject), ('macro_result', macro_result), ('micro_result', micro_result)]:
            if value:
                clauses.append(f'{field} = ?')
                params.append(value)
        if subsidy:
            clauses.append(f'{subsidy} = ?')
            params.append(1 if presence == 'provided' else 0)
        where = (' WHERE ' + ' AND '.join(clauses)) if clauses else ''
        with self._lock:
            self._load()
            total = self._db.execute('SELECT count(*) FROM history' + where, params).fetchone()[0]
            pages = max(1, (total + page_size - 1) // page_size)
            page = min(page, pages)
            rows = self._db.execute('SELECT * FROM history' + where +
                f' ORDER BY {sort} {direction}, source_row ASC LIMIT ? OFFSET ?',
                [*params, page_size, (page - 1) * page_size]).fetchall()
            items = [{key: row[key] for key in row.keys() if key != 'search_text'} for row in rows]
            return dict(items=items, total=total, total_cases=self._total, page=page, page_size=page_size,
                        total_pages=pages, options=self._options, source=self.source.name,
                        result_sheet=RESULT_SHEET, subsidy_sheet=SUBSIDY_SHEET)
