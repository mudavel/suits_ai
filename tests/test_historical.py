from collections import Counter
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.backend.config import Settings
from src.backend.main import create_app
from src.backend.services.historical_service import HistoricalBase, RESULT_SHEET, SUBSIDY_SHEET

SOURCE = Path(__file__).resolve().parents[1] / 'artefacts/Hackaton Unicamp/Hackaton_Enter_Base_Candidatos.xlsx'


@pytest.fixture(scope='module')
def historical():
    base = HistoricalBase(SOURCE)
    yield base
    base.close()


def test_all_60000_cases_match_both_excel_sheets(historical):
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    try:
        results = list(workbook[RESULT_SHEET].iter_rows(min_row=2, values_only=True))
        subsidies = {row[0]: (number, row[1:]) for number, row in enumerate(workbook[SUBSIDY_SHEET].iter_rows(min_row=3, values_only=True), 3)}
    finally:
        workbook.close()
    assert len(results) == len(subsidies) == 60000
    actual = []
    for page in range(1, 601):
        response = historical.query(page=page, page_size=100)
        assert response['total'] == response['total_cases'] == 60000
        assert response['total_pages'] == 600
        actual.extend(response['items'])
    assert len({item['case_number'] for item in actual}) == 60000
    for source_row, (item, row) in enumerate(zip(actual, results, strict=True), 2):
        assert tuple(item[field] for field in ('case_number', 'uf', 'subject', 'sub_subject', 'macro_result', 'micro_result', 'cause_value', 'condemnation_value')) == row
        subsidy_row, values = subsidies[item['case_number']]
        assert tuple(item[field] for field in ('contract', 'statement', 'credit_receipt', 'dossier', 'debt_evolution', 'referenced_report')) == values
        assert item['source_row'] == source_row
        assert item['subsidy_row'] == subsidy_row
        assert item['subsidy_count'] == sum(values)
    assert Counter(item['micro_result'] for item in actual)['Acordo'] == 280


def test_global_search_filters_numeric_sort_and_page_boundaries(historical):
    last = historical.query(page=600, page_size=100)['items'][-1]
    found = historical.query(q=last['case_number'])
    assert found['total'] == 1 and found['items'][0] == last
    filters = dict(uf=last['uf'], micro_result=last['micro_result'], subsidy='contract', presence='provided' if last['contract'] else 'missing')
    matches = historical.query(**filters, sort='condemnation_value', direction='desc', page_size=100)
    assert matches['total'] > 0
    assert all(item['uf'] == last['uf'] and item['micro_result'] == last['micro_result'] and item['contract'] == last['contract'] for item in matches['items'])
    values = [item['condemnation_value'] for item in matches['items']]
    assert values == sorted(values, reverse=True)
    macro = next(value for value in matches['options']['macro_result'] if value.casefold() == 'não êxito')
    assert historical.query(q='nao exito')['total'] == historical.query(macro_result=macro)['total']
    assert historical.query(q="' OR 1=1 --")['total'] == 0
    assert historical.query(q='processo-inexistente', page=100)['page'] == 1
    assert historical.query(page=99999)['page'] == 1200
    assert historical.query(q=last['case_number'], subsidy='contract', presence='missing' if last['contract'] else 'provided')['total'] == 0


def test_endpoint_validation_and_source_failure(tmp_path, historical):
    settings = Settings(data_mode='mock', database_path=tmp_path / 'api.sqlite3', artifacts_dir=SOURCE.parent)
    with TestClient(create_app(settings)) as client:
        client.app.state.historical = historical
        response = client.get('/api/monitoring/historical', params={'page_size': 25, 'micro_result': 'Acordo'})
        assert response.status_code == 200
        data = response.json()
        assert data['total'] == 280 and len(data['items']) == 25
        assert data['source'] == SOURCE.name
        for params in ({'page': 0}, {'page_size': 101}, {'sort': 'search_text'}, {'direction': 'invalid'}, {'subsidy': 'unknown'}):
            assert client.get('/api/monitoring/historical', params=params).status_code == 422
        client.app.state.historical = HistoricalBase(tmp_path / 'missing.xlsx')
        assert client.get('/api/monitoring/historical').status_code == 503
