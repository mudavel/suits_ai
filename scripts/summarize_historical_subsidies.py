"""Print the historical document inventory as JSON; never change the source workbook.

The dashboard snapshot is src/frontend/src/data/historicalSubsidies.json.
Run this script when updating the supplied historical workbook.
"""
import json
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(__file__).resolve().parents[1] / "artefacts/Hackaton Unicamp/Hackaton_Enter_Base_Candidatos.xlsx"
COLUMNS = {
    "Contrato": "has_contract",
    "Extrato": "has_statement",
    "Comprovante de crédito": "has_credit_receipt",
    "Dossiê": "has_dossier",
    "Demonstrativo de evolução da dívida": "has_debt_evolution",
    "Laudo referenciado": "has_referenced_report",
}


def summarize(source=SOURCE):
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        results = workbook["Resultados dos processos"].iter_rows(values_only=True)
        result_headers = next(results)
        result_key = result_headers.index("Número do processo")
        result_ids = [row[result_key] for row in results if row[result_key] is not None]
        rows = workbook["Subsídios disponibilizados"].iter_rows(min_row=2, values_only=True)
        headers = next(rows)
        key = headers.index("Número do processos")
        positions = {headers.index(label): field for label, field in COLUMNS.items()}
        missing = dict.fromkeys(COLUMNS.values(), 0)
        seen = set()
        for row in rows:
            if row[key] is None or row[key] in seen:
                raise ValueError("Identificador de subsídio ausente ou duplicado.")
            seen.add(row[key])
            for index, field in positions.items():
                if row[index] not in (0, 1):
                    raise ValueError(f"Presença documental inválida: {headers[index]}")
                missing[field] += int(row[index] == 0)
        if len(result_ids) != len(set(result_ids)) or set(result_ids) != seen or len(seen) != 60000:
            raise ValueError("A base deve conter os mesmos 60.000 processos únicos nas duas abas.")
        return {
            "total_cases": len(seen),
            "missing_by_type": missing,
            "source": source.name,
            "sheet": "Subsídios disponibilizados",
            "range": "A2:G60002",
            "method": "Contagem de valores 0 (subsídio não fornecido), por processo único; 1 indica fornecido.",
        }
    finally:
        workbook.close()


if __name__ == "__main__":
    print(json.dumps(summarize(), ensure_ascii=False, indent=2))
