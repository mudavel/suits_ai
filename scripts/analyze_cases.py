"""Read-only exploratory analysis of the hackathon workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from render_analysis import render


DOCUMENTS = ["contrato", "extrato", "comprovante_credito", "dossie", "demonstrativo_divida", "laudo"]
LABELS = {"contrato": "Contrato", "extrato": "Extrato", "comprovante_credito": "Comprovante de crédito", "dossie": "Dossiê", "demonstrativo_divida": "Demonstrativo da dívida", "laudo": "Laudo referenciado"}
LOSS_OUTCOMES = ["Procedência", "Parcial procedência"]
UF_CODES = set("AC AL AP AM BA CE DF ES GO MA MT MS MG PA PB PR PE PI RJ RN RS RO RR SC SP SE TO".split())


def load_cases(path: Path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = workbook["Resultados dos processos"].iter_rows(values_only=True)
    result_headers = list(next(rows))
    assert result_headers == ["Número do processo", "UF", "Assunto", "Sub-assunto", "Resultado macro", "Resultado micro", "Valor da causa", "Valor da condenação/indenização"], "Unexpected results schema."
    results = pd.DataFrame(rows, columns=["processo", "uf", "assunto", "subassunto", "resultado_macro", "resultado_micro", "valor_causa", "valor_resultado"])
    rows = workbook["Subsídios disponibilizados"].iter_rows(values_only=True)
    legend = list(next(rows))
    subsidy_headers = list(next(rows))
    assert subsidy_headers == ["Número do processos", "Contrato", "Extrato", "Comprovante de crédito", "Dossiê", "Demonstrativo de evolução da dívida", "Laudo referenciado"], "Unexpected document schema."
    subsidies = pd.DataFrame(rows, columns=["processo", *DOCUMENTS])
    workbook.close()
    quality = {
        "result_rows": len(results),
        "subsidy_rows": len(subsidies),
        "result_headers": result_headers,
        "subsidy_headers": subsidy_headers,
        "legend": legend,
        "result_missing": results.isna().sum().to_dict(),
        "subsidy_missing": subsidies.isna().sum().to_dict(),
        "duplicate_result_ids": int(results.processo.duplicated().sum()),
        "duplicate_subsidy_ids": int(subsidies.processo.duplicated().sum()),
        "result_ids_without_subsidies": len(set(results.processo) - set(subsidies.processo)),
        "subsidy_ids_without_results": len(set(subsidies.processo) - set(results.processo)),
        "document_values": {column: subsidies[column].value_counts(dropna=False).to_dict() for column in DOCUMENTS},
    }
    cases = results.merge(subsidies, on="processo", how="outer", validate="one_to_one", indicator=True)
    assert cases._merge.eq("both").all(), "The two sheets do not contain the same cases."
    assert not cases.drop(columns="_merge").isna().any().any(), "Missing input values."
    assert cases.processo.map(lambda value: isinstance(value, str) and bool(value.strip())).all(), "Invalid process identifiers."
    assert cases.resultado_micro.isin(["Procedência", "Parcial procedência", "Improcedência", "Extinção", "Acordo"]).all(), "Unknown outcome."
    assert np.isfinite(cases[["valor_causa", "valor_resultado"]].to_numpy(float)).all(), "Nonfinite amounts."
    assert cases.valor_causa.gt(0).all() and cases.valor_resultado.ge(0).all(), "Invalid amounts."
    assert all(cases[column].isin([0, 1]).all() for column in DOCUMENTS)
    cases[DOCUMENTS] = cases[DOCUMENTS].astype(int)
    cases["document_count"] = cases[DOCUMENTS].sum(axis=1)
    return cases, quality


def records(frame):
    return json.loads(frame.to_json(orient="records", force_ascii=False))


def wilson(successes, n):
    if not n:
        return None, None
    z = 1.959963984540054
    rate = successes / n
    middle = (rate + z * z / (2 * n)) / (1 + z * z / n)
    half = z * math.sqrt(rate * (1 - rate) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return middle - half, middle + half


def summarize_group(group):
    court = group.loc[~group.is_agreement]
    losses = group.loc[group.is_loss]
    lo, hi = wilson(len(losses), len(court))
    return {
        "n": len(group), "n_court": len(court), "losses": len(losses),
        "agreements": int(group.is_agreement.sum()),
        "loss_rate": len(losses) / len(court) if len(court) else None,
        "loss_rate_ci_low": lo, "loss_rate_ci_high": hi,
        "cost_mean_all": float(group.valor_resultado.mean()),
        "cost_mean_court": float(court.valor_resultado.mean()) if len(court) else None,
        "cost_mean_loss": float(losses.valor_resultado.mean()) if len(losses) else None,
        "cost_median_loss": float(losses.valor_resultado.median()) if len(losses) else None,
        "cost_p90_loss": float(losses.valor_resultado.quantile(.9)) if len(losses) else None,
        "claim_mean": float(group.valor_causa.mean()),
        "total_result": float(group.valor_resultado.sum()),
        "payout_ratio_mean_loss": float(losses.payout_ratio.mean()) if len(losses) else None,
    }


def group_table(cases, keys):
    keys = keys if isinstance(keys, list) else [keys]
    output = []
    for name, group in cases.groupby(keys, observed=True, sort=True):
        values = name if isinstance(name, tuple) else (name,)
        output.append({**dict(zip(keys, values)), **summarize_group(group)})
    return pd.DataFrame(output)


def sigmoid(x):
    return 1 / (1 + np.exp(-np.clip(x, -35, 35)))


def fit_logistic(x, y, penalty=1.0):
    coefficients = np.zeros(x.shape[1])
    regularization = np.eye(x.shape[1]) * penalty
    regularization[0, 0] = 0
    for iteration in range(40):
        p = sigmoid(x @ coefficients)
        weights = np.maximum(p * (1 - p), 1e-8)
        hessian = x.T @ (x * weights[:, None]) + regularization
        gradient = x.T @ (y - p) - regularization @ coefficients
        step = np.linalg.solve(hessian, gradient)
        coefficients += step
        if np.max(np.abs(step)) < 1e-7:
            break
    return coefficients, iteration + 1


def feature_matrix(frame, mode, training=None):
    config = {} if training is None else training.copy()
    columns = [np.ones(len(frame))]
    names = ["intercept"]
    if mode in ["uf", "controls", "full"]:
        levels = config.setdefault("uf_levels", sorted(frame.uf.unique()))
        for level in levels[1:]:
            columns.append(frame.uf.eq(level).to_numpy(float))
            names.append("uf_" + level)
    if mode in ["controls", "full"]:
        columns.append(frame.subassunto.eq("Golpe").to_numpy(float))
        names.append("subassunto_golpe")
        log_claim = np.log(frame.valor_causa.to_numpy(float))
        config.setdefault("log_claim_mean", float(log_claim.mean()))
        config.setdefault("log_claim_std", float(log_claim.std()))
        columns.append((log_claim - config["log_claim_mean"]) / config["log_claim_std"])
        names.append("log_valor_causa_padronizado")
    if mode in ["documents", "full"]:
        for document in DOCUMENTS:
            columns.append(frame[document].to_numpy(float))
            names.append(document)
    return np.column_stack(columns), names, config


def classification_metrics(y, p):
    p = np.clip(p, 1e-12, 1 - 1e-12)
    ranks = pd.Series(p).rank(method="average").to_numpy()
    positive = y == 1
    n_positive = positive.sum()
    n_negative = (~positive).sum()
    auc = (ranks[positive].sum() - n_positive * (n_positive + 1) / 2) / (n_positive * n_negative)
    n_selected = round(len(p) * .2)
    threshold = np.sort(p)[-n_selected]
    above, tied = p > threshold, p == threshold
    # Fractional allocation at the cutoff prevents label/order-dependent tie breaking.
    expected_selected_losses = y[above].sum() + (n_selected - above.sum()) * y[tied].mean()
    return {"auc": float(auc), "brier": float(np.mean((p - y) ** 2)), "log_loss": float(-np.mean(y * np.log(p) + (1 - y) * np.log(1 - p))), "observed_rate": float(y.mean()), "predicted_rate": float(p.mean()), "top_20pct_loss_rate": float(expected_selected_losses / n_selected), "top_20pct_loss_capture": float(expected_selected_losses / y.sum())}


def run_holdout(cases):
    court = cases.loc[~cases.is_agreement].copy().reset_index(drop=True)
    rng = np.random.default_rng(20260912)
    train_indices, test_indices = [], []
    for target in [False, True]:
        indices = rng.permutation(np.flatnonzero(court.is_loss.to_numpy() == target))
        split = int(.8 * len(indices))
        train_indices.extend(indices[:split])
        test_indices.extend(indices[split:])
    train = court.iloc[rng.permutation(train_indices)].copy()
    test = court.iloc[rng.permutation(test_indices)].copy()
    y_train = train.is_loss.to_numpy(float)
    y_test = test.is_loss.to_numpy(float)
    outputs = []
    predictions = {}
    models = {}
    for mode in ["global", "uf", "documents", "controls", "full"]:
        x_train, names, config = feature_matrix(train, mode)
        x_test, _, _ = feature_matrix(test, mode, config)
        coefficients, iterations = fit_logistic(x_train, y_train)
        p = sigmoid(x_test @ coefficients)
        predictions[mode] = p
        outputs.append({"model": mode, "features": len(names) - 1, "iterations": iterations, **classification_metrics(y_test, p)})
        models[mode] = (coefficients, config)
    # Severity uses only the training cases with judicial condemnation.
    # Fit the paid/claimed ratio, with the same pre-outcome features, then bound to [0, 1].
    x_train, names, config = feature_matrix(train, "full")
    x_test, _, _ = feature_matrix(test, "full", config)
    loss_mask = y_train == 1
    severity_penalty = np.eye(x_train.shape[1])
    severity_penalty[0, 0] = 0
    severity = np.linalg.solve(x_train[loss_mask].T @ x_train[loss_mask] + severity_penalty, x_train[loss_mask].T @ train.loc[train.is_loss, "payout_ratio"].to_numpy(float))
    cost_predictions = {
        "always_zero": np.zeros(len(test)),
        "global_mean": np.full(len(test), train.valor_resultado.mean()),
        "probability_times_mean_severity": predictions["full"] * train.loc[train.is_loss, "valor_resultado"].mean(),
        "two_part_claim_ratio": predictions["full"] * np.clip(x_test @ severity, 0, 1) * test.valor_causa.to_numpy(float),
    }
    cost_metrics = []
    actual = test.valor_resultado.to_numpy(float)
    for name, estimate in cost_predictions.items():
        cost_metrics.append({"model": name, "mae": float(np.abs(actual - estimate).mean()), "rmse": float(np.sqrt(np.mean((actual - estimate) ** 2))), "observed_mean": float(actual.mean()), "predicted_mean": float(estimate.mean()), "relative_total_error": float(estimate.sum() / actual.sum() - 1)})
    test["predicted_loss"] = predictions["full"]
    test["predicted_cost"] = cost_predictions["two_part_claim_ratio"]
    test["risk_decile"] = pd.qcut(test.predicted_loss.rank(method="first"), 10, labels=False) + 1
    calibration = test.groupby("risk_decile", observed=True).agg(n=("processo", "size"), predicted_loss=("predicted_loss", "mean"), observed_loss=("is_loss", "mean"), predicted_cost=("predicted_cost", "mean"), observed_cost=("valor_resultado", "mean")).reset_index()
    priority = test.nlargest(round(.2 * len(test)), "predicted_loss")
    priority_summary = {"n": len(priority), "loss_rate": float(priority.is_loss.mean()), "loss_capture": float(priority.is_loss.sum() / test.is_loss.sum()), "observed_cost_share": float(priority.valor_resultado.sum() / test.valor_resultado.sum()), "min_predicted_loss": float(priority.predicted_loss.min())}
    bootstrap_auc = []
    for _ in range(200):
        sample = rng.integers(0, len(test), size=len(test))
        bootstrap_auc.append(classification_metrics(y_test[sample], predictions["full"][sample])["auc"])
    document_effects = []
    x_train, names, config = feature_matrix(train, "full")
    coef = models["full"][0]
    for document in DOCUMENTS:
        absent = test.copy()
        present = test.copy()
        absent[document] = 0
        present[document] = 1
        x_absent, _, _ = feature_matrix(absent, "full", config)
        x_present, _, _ = feature_matrix(present, "full", config)
        # Standardized model association; changing a feature is not a causal intervention.
        document_effects.append({"document": document, "odds_ratio_conditional": float(np.exp(coef[names.index(document)])), "standardized_predicted_difference": float((sigmoid(x_present @ coef) - sigmoid(x_absent @ coef)).mean())})
    full_x_train, full_names, _ = feature_matrix(train, "full")
    gradient = full_x_train.T @ (y_train - sigmoid(full_x_train @ coef))
    gradient[1:] -= coef[1:]
    assert float(np.abs(gradient).max()) < 1e-4, "Logistic regression did not converge."
    return {"seed": 20260912, "target": "Parcial procedência ou Procedência entre processos encerrados sem acordo; extinções são mantidas como não condenação.", "train_n": len(train), "test_n": len(test), "excluded_agreements": int(cases.is_agreement.sum()), "classification": outputs, "cost_prediction": cost_metrics, "calibration": records(calibration), "conditional_document_associations": document_effects, "full_auc_bootstrap_ci95": np.quantile(bootstrap_auc, [.025, .975]).tolist(), "bootstrap_replicates": len(bootstrap_auc), "top_20pct_priority": priority_summary, "full_model_coefficients": dict(zip(full_names, map(float, coef))), "max_abs_penalized_gradient": float(np.abs(gradient).max())}


def analyze(cases, quality):
    cases["is_agreement"] = cases.resultado_micro.eq("Acordo")
    cases["is_loss"] = cases.resultado_micro.isin(LOSS_OUTCOMES)
    cases["payout_ratio"] = cases.valor_resultado / cases.valor_causa
    cases["claim_band"] = pd.cut(cases.valor_causa, [0, 5000, 10000, 15000, 20000, 25000, np.inf], right=False, labels=["<5 mil", "5–10 mil", "10–15 mil", "15–20 mil", "20–25 mil", "25 mil+"])
    cases["document_pattern"] = cases[DOCUMENTS].astype(str).agg("".join, axis=1)
    tables = {name: group_table(cases, keys) for name, keys in {
        "by_uf": "uf", "by_subsubject": "subassunto", "by_document_count": "document_count", "by_claim_band": "claim_band", "by_contract_statement": ["contrato", "extrato"], "by_core_documents": ["contrato", "extrato", "comprovante_credito"], "by_document_pattern": "document_pattern", "by_uf_subsubject": ["uf", "subassunto"],
    }.items()}
    documents = []
    for document in DOCUMENTS:
        absent = cases.loc[cases[document] == 0]
        present = cases.loc[cases[document] == 1]
        a, p = summarize_group(absent), summarize_group(present)
        documents.append({"document": document, "label": LABELS[document], "available_n": len(present), "availability": len(present) / len(cases), "absent_n_court": a["n_court"], "present_n_court": p["n_court"], "absent_loss_rate": a["loss_rate"], "present_loss_rate": p["loss_rate"], "difference_present_minus_absent": p["loss_rate"] - a["loss_rate"], "absent_cost_mean_court": a["cost_mean_court"], "present_cost_mean_court": p["cost_mean_court"]})
    tables["by_document"] = pd.DataFrame(documents)
    merits = cases.loc[cases.resultado_micro.isin([*LOSS_OUTCOMES, "Improcedência"])]
    tables["merits_by_document_count"] = group_table(merits, "document_count")
    tables["merits_by_contract_statement"] = group_table(merits, ["contrato", "extrato"])
    quality.update({"uf_count": int(cases.uf.nunique()), "missing_ufs": sorted(UF_CODES - set(cases.uf)), "unknown_ufs": sorted(set(cases.uf) - UF_CODES), "uf_min_n": int(cases.uf.value_counts().min()), "uf_max_n": int(cases.uf.value_counts().max()), "constant_columns": [column for column in ["assunto", "subassunto", *DOCUMENTS] if cases[column].nunique() == 1], "macro_micro_mapping_consistent": bool((cases.resultado_macro.eq("Não Êxito") == (cases.is_loss | cases.is_agreement)).all()), "money_outcome_consistent": bool((cases.valor_resultado.gt(0) == (cases.is_loss | cases.is_agreement)).all()), "results_more_than_two_decimals": int((np.abs(cases.valor_resultado * 100 - np.round(cases.valor_resultado * 100)) > 1e-6).sum())})
    ratios = cases.loc[cases.is_loss].groupby("resultado_micro").payout_ratio.agg(["count", "mean", "median", "min", "max"]).reset_index()
    ratio_frequencies = {outcome: {str(key): int(n) for key, n in group.payout_ratio.round(6).value_counts().head(12).items()} for outcome, group in cases.loc[cases.valor_resultado.gt(0)].groupby("resultado_micro")}
    agreements = cases.loc[cases.is_agreement]
    full_cost = float(cases.valor_resultado.sum())
    # This is a descriptive concentration, not savings or a policy backtest.
    sorted_cost = cases.valor_resultado.sort_values(ascending=False)
    extended = {"overall": summarize_group(cases), "merits": {"n": len(merits), "losses": int(merits.is_loss.sum()), "loss_rate": float(merits.is_loss.mean())}, "agreements": {"n": len(agreements), "mean": float(agreements.valor_resultado.mean()), "median": float(agreements.valor_resultado.median()), "mean_claim_ratio": float(agreements.payout_ratio.mean()), "uf_counts": agreements.uf.value_counts().to_dict()}, "ratios": records(ratios), "ratio_frequencies": ratio_frequencies, "cost_concentration_top_10pct": float(sorted_cost.head(round(.1 * len(cases))).sum() / full_cost), "tables": {name: records(table) for name, table in tables.items()}, "validation": run_holdout(cases)}
    return extended, tables


def verify_math():
    y = np.array([0., 0., 1., 1.])
    assert classification_metrics(y, np.array([.1, .2, .8, .9]))["auc"] == 1
    assert classification_metrics(y, np.array([.9, .8, .2, .1]))["auc"] == 0
    tied = classification_metrics(y, np.full(4, .5))
    assert tied["auc"] == .5 and tied["top_20pct_loss_rate"] == .5
    coefficients, _ = fit_logistic(np.ones((10, 1)), np.array([0.] * 7 + [1.] * 3))
    assert abs(sigmoid(coefficients)[0] - .3) < 1e-8


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parents[1] / "artefacts/suits_docs/analysis")
    args = parser.parse_args()
    verify_math()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    cases, quality = load_cases(args.workbook)
    extended, tables = analyze(cases, quality)
    profile = {
        "source": args.workbook.name,
        "sha256": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
        "quality": quality,
        "categories": {column: cases[column].value_counts().to_dict() for column in ["uf", "assunto", "subassunto", "resultado_macro", "resultado_micro", "document_count"]},
        "macro_micro": records(cases.groupby(["resultado_macro", "resultado_micro"], dropna=False).agg(n=("processo", "size"), total=("valor_resultado", "sum"), mean=("valor_resultado", "mean"), median=("valor_resultado", "median"), positive=("valor_resultado", lambda x: int(x.gt(0).sum())), minimum=("valor_resultado", "min"), maximum=("valor_resultado", "max")).reset_index()),
        "money_summary": records(cases[["valor_causa", "valor_resultado"]].describe(percentiles=[.01, .05, .25, .5, .75, .9, .95, .99]).reset_index()),
        "money_quality": {"negative_cause": int(cases.valor_causa.lt(0).sum()), "negative_result": int(cases.valor_resultado.lt(0).sum()), "zero_cause": int(cases.valor_causa.eq(0).sum()), "zero_result": int(cases.valor_resultado.eq(0).sum()), "result_exceeds_cause": int(cases.valor_resultado.gt(cases.valor_causa).sum()), "nonfinite_cause": int((~np.isfinite(cases.valor_causa)).sum()), "nonfinite_result": int((~np.isfinite(cases.valor_resultado)).sum())},
        **extended,
    }
    for name, table in tables.items():
        table.to_csv(args.output_dir / (name + ".csv"), index=False, encoding="utf-8-sig")
    (args.output_dir / "profile.json").write_text(json.dumps(profile, ensure_ascii=False, indent=2, default=int), encoding="utf-8")
    render(profile, args.output_dir)
    console = {"output_dir": str(args.output_dir), "overall": profile["overall"], "classification": profile["validation"]["classification"], "auc_ci95": profile["validation"]["full_auc_bootstrap_ci95"], "source_sha256": profile["sha256"]}
    print(json.dumps(console, ensure_ascii=True, indent=2, default=int))


if __name__ == "__main__":
    main()
